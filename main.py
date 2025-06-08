import streamlit as st
import tempfile
import os
import mammoth
from docx import Document
import PyPDF2
import re
import google.generativeai as genai
from google.api_core import exceptions as api_exceptions
from bs4 import BeautifulSoup
import requests
import nltk
from langdetect import detect
from collections import OrderedDict
import json
from openai import OpenAI
import datetime

# --- ReportLab Imports for PDF generation ---
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT, TA_JUSTIFY
from reportlab.lib.colors import black, blue, darkblue
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Image, ListFlowable, ListItem, Table, TableStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus.flowables import Flowable
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from reportlab.graphics.shapes import Drawing, Line
from reportlab.graphics import renderPDF

# --- Ensure NLTK stopwords are downloaded ---
try:
    nltk.data.find('corpora/stopwords')
except nltk.downloader.DownloadError:
    nltk.download('stopwords', quiet=True)
from nltk.corpus import stopwords

# --- Utility Functions ---

def extract_text(file_path):
    """Extracts text from DOCX or PDF files."""
    if file_path.endswith('.docx'):
        doc = Document(file_path)
        return '\n'.join([p.text.strip() for p in doc.paragraphs if p.text.strip()])
    elif file_path.endswith('.pdf'):
        text = ""
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() + "\n"
        return text
    else:
        st.warning(f"File type {file_path.split('.')[-1]} not directly supported for text extraction.")
        return ""

def detect_language(text):
    """Detects language of the text, defaults to English if detection fails."""
    try:
        lang = detect(text)
        return 'french' if lang == 'fr' else 'english'
    except:
        return 'english'

def fetch_job_description(input_text_or_url):
    """Fetches job description from URL or returns the input text directly."""
    if input_text_or_url.strip().lower().startswith('http'):
        try:
            resp = requests.get(input_text_or_url, timeout=10)
            resp.raise_for_status()
            soup = BeautifulSoup(resp.text, 'html.parser')
            for script_or_style in soup(["script", "style"]):
                script_or_style.extract()
            text = soup.get_text(separator=' ', strip=True)
            return text
        except requests.exceptions.RequestException as e:
            st.error(f"Error fetching URL: {e}")
            return ""
    else:
        return input_text_or_url

def extract_keywords(text, top_n=15, language='english'):
    """Extracts top N keywords from text, excluding stopwords."""
    stop_words = set(stopwords.words(language))
    words = re.findall(r'\b\w+\b', text.lower())
    freq = {}
    for word in words:
        if word not in stop_words and len(word) > 2:
            freq[word] = freq.get(word, 0) + 1
    sorted_keywords = sorted(freq, key=freq.get, reverse=True)
    return sorted_keywords[:top_n]

def extract_personal_details(text):
    """Extracts personal information from resume text."""
    personal_info = {
        "name": "",
        "email": "",
        "phone": "",
        "linkedin": "",
        "github": "",
        "address": "",
        "city": ""
    }

    lines = text.splitlines()[:20]  # Check first 20 lines

    # Extract name (first substantial line that doesn't contain contact info)
    for line in lines:
        line = line.strip()
        if not line:
            continue

        if not re.search(r'(@|\.com|github|linkedin|http|https|\d[\d\s\-()]+\d)', line, re.IGNORECASE) and len(line.split()) <= 4:
            personal_info["name"] = line
            break

    full_text = '\n'.join(lines)

    # Email
    email_match = re.search(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', full_text)
    if email_match:
        personal_info["email"] = email_match.group(0)

    # Phone
    phone_match = re.search(r'(\+?\d{1,3}[\s-]?)?(\(?\d{3}\)?[\s-]?\d{3}[\s-]?\d{4}|\d{10})', full_text)
    if phone_match:
        personal_info["phone"] = phone_match.group(0)

    # LinkedIn
    linkedin_match = re.search(r'(linkedin\.com/in/[a-zA-Z0-9_-]+)', full_text, re.IGNORECASE)
    if linkedin_match:
        personal_info["linkedin"] = "https://" + linkedin_match.group(0)

    # GitHub
    github_match = re.search(r'(github\.com/[a-zA-Z0-9_-]+)', full_text, re.IGNORECASE)
    if github_match:
        personal_info["github"] = "https://" + github_match.group(0)

    # Address and city (basic extraction)
    address_patterns = [
        r'\d+\s+[A-Za-z\s]+(?:Street|St|Avenue|Ave|Road|Rd|Drive|Dr|Lane|Ln)',
        r'[A-Za-z\s]+,\s*[A-Z]{2}\s*\d{5}'
    ]
    for pattern in address_patterns:
        match = re.search(pattern, full_text, re.IGNORECASE)
        if match:
            if not personal_info["address"]:
                personal_info["address"] = match.group(0)
            break

    # Attempt to extract city from address or general text
    city_match = re.search(r'([A-Z][a-z]+(?: [A-Z][a-z]+)*),\s*[A-Z]{2}', full_text)
    if city_match:
        personal_info["city"] = city_match.group(1)

    return personal_info

def parse_resume_sections(text):
    """Parses resume text into sections, with special handling for Experience and Projects."""
    sections = OrderedDict()
    current_section = None

    common_headings = [
        "professional summary", "summary", "objective", "about me", "profile",
        "experience", "work experience", "professional experience", "employment",
        "projects", "portfolio",
        "education", "academic background",
        "skills", "technical skills", "core competencies", "expertise", "competencies",
        "certifications", "awards", "achievements",
        "publications", "volunteer", "volunteering",
        "languages", "hobbies", "interests", "references",
        "contact"
    ]

    lines = text.split('\n')
    temp_content_buffer = [] # Buffer to hold lines before they are assigned to a section

    def _process_buffer(section_name, buffer):
        if not buffer:
            return []

        # Special handling for Experience and Projects to structure them
        if section_name in ["Experience", "Projects"]:
            parsed_entries = []
            current_entry = None

            # Regex to identify potential titles for Experience/Projects
            # This is a heuristic and might need fine-tuning based on common resume formats
            # Examples:
            # - Company – Role | Dates (e.g., Infosys – Senior System Engineer | Jan 2024 - Present)
            # - Project Name (e.g., Journal App)
            # - Any line that looks like a strong start to an entry (starts with capitalized word, contains keywords)
            entry_title_patterns = [
                re.compile(r'^[A-Z][a-zA-Z\s,]+\s*–\s*[A-Z][a-zA-Z\s,]+\s*\|\s*\w{3}\s*\d{4}\s*-\s*(?:\w{3}\s*\d{4}|Present)'),
                re.compile(r'^[A-Z][a-zA-Z\s]+\s*(?:Engineer|Developer|Manager|Analyst|Specialist)'), # Role + Title
                re.compile(r'^(?:<b>)?[A-Za-z][a-zA-Z\s]+(?:App|Platform|System|Tool)(?:</b>)?'), # Project Name type
                re.compile(r'^[A-Z][a-zA-Z\s]+(?:\s*\d{4})?') # General capitalized line, potentially a title
            ]
            
            # Helper to check if a line is a likely bullet point
            def is_likely_bullet(line_text):
                # Check for common action verbs or sentence structure of a bullet
                action_verbs = ["Designed", "Developed", "Led", "Implemented", "Managed", "Built", "Optimized", "Achieved", "Spearheaded", "Improved", "Diagnosed"]
                for verb in action_verbs:
                    if line_text.strip().lower().startswith(verb.lower()):
                        return True
                # Check for quantification (e.g., "by 40%")
                if re.search(r'\d+%|\$\d+|[xX]\d+', line_text):
                    return True
                return False

            for i, line in enumerate(buffer):
                line_stripped = line.strip()
                if not line_stripped:
                    continue

                is_new_entry_title = False
                for pattern in entry_title_patterns:
                    if pattern.match(line_stripped) and len(line_stripped.split()) < 15: # Prevent very long lines from being titles
                        is_new_entry_title = True
                        break
                
                # Also consider lines that look like main company/role/project names that aren't action verbs
                # and are followed by what appears to be a bullet or sub-details
                if not is_new_entry_title and current_entry and not is_likely_bullet(line_stripped) and len(line_stripped.split()) < 10:
                    # If current line is short, not a bullet, and we have an entry, it might be a subtitle for the current entry
                    # Append it to the title of the current entry or keep it in bullets if it looks like a sub-detail
                    current_entry["bullets"].append(line_stripped) # Treat as a descriptive bullet for now
                    continue


                if is_new_entry_title:
                    if current_entry:
                        parsed_entries.append(current_entry)
                    current_entry = {"title": line_stripped, "bullets": []}
                elif current_entry:
                    current_entry["bullets"].append(line_stripped)
                else: # Content before the first detected entry, or if no title is ever found
                    if not parsed_entries: # If no entries yet, create a dummy one or append to a conceptual 'header'
                        parsed_entries.append({"title": "", "bullets": [line_stripped]})
                    else: # If previous entries exist but this line doesn't fit a new title, append to last entry's bullets
                         parsed_entries[-1]["bullets"].append(line_stripped)

            if current_entry:
                parsed_entries.append(current_entry)

            # Final cleanup: remove entries with empty titles and no bullets, or empty bullets
            final_entries = []
            for entry in parsed_entries:
                if entry.get("title") or entry.get("bullets"): # Keep if has title OR bullets
                    if entry.get("bullets"):
                        entry["bullets"] = [b for b in entry["bullets"] if b.strip()] # Clean empty bullets
                    final_entries.append(entry)

            # Fallback: if structured parsing yields nothing, return original buffer as simple strings
            if not final_entries and buffer:
                 return [line.strip() for line in buffer if line.strip()]

            return final_entries if final_entries else []
        else:
            # For other sections, return simple list of non-empty strings
            return [line.strip() for line in buffer if line.strip()]


    for line in lines:
        line = line.strip()
        if not line:
            continue

        is_heading_found = False
        potential_heading_lower = line.replace(':', '').strip().lower()

        # Check for exact matches or very close matches first
        for ch in common_headings:
            if ch == potential_heading_lower:
                if current_section: # If we have an active section, process its buffer
                    sections[current_section] = _process_buffer(current_section, temp_content_buffer)
                    temp_content_buffer = [] # Reset buffer
                current_section = ch.title()
                is_heading_found = True
                break
            # Allow for partial matches, but heading shouldn't be too long and should be a whole word
            if re.search(r'\b' + re.escape(ch) + r'\b', potential_heading_lower) and len(potential_heading_lower.split()) < 5:
                if current_section:
                    sections[current_section] = _process_buffer(current_section, temp_content_buffer)
                    temp_content_buffer = []
                current_section = ch.title() # Use the common heading name
                is_heading_found = True
                break

        # Fallback heading detection (e.g., ALL CAPS lines, or lines ending with colon)
        if not is_heading_found and ((line.isupper() and len(line.split()) < 6) or \
           (re.match(r'^[A-Za-z\s&]+:\s*$', line) and len(line.split()) < 6)):
            if current_section:
                sections[current_section] = _process_buffer(current_section, temp_content_buffer)
                temp_content_buffer = []
            current_section = line.rstrip(':').strip().title()
            is_heading_found = True
        
        # Standardize section names after detection
        if is_heading_found:
            if current_section.lower() in ["work experience", "professional experience"]:
                current_section = "Experience"
            elif current_section.lower() in ["technical skills", "core competencies", "expertise", "competencies"]:
                current_section = "Skills"
            elif current_section.lower() in ["academic background"]:
                current_section = "Education"
            elif current_section.lower() in ["professional summary", "objective", "about me", "profile"]:
                current_section = "Professional Summary"
            
            if current_section not in sections:
                sections[current_section] = [] # Initialize if new section

        elif current_section:
            temp_content_buffer.append(line)
        else: # Lines before any recognized section (often part of an implicit header/summary)
            if "Header/Summary" not in sections:
                sections["Header/Summary"] = []
            sections["Header/Summary"].append(line)

    # Process any remaining content in the buffer after the loop finishes
    if current_section and temp_content_buffer:
        sections[current_section] = _process_buffer(current_section, temp_content_buffer)

    # Clean empty sections
    cleaned_sections = OrderedDict()
    for k, v in sections.items():
        if v: # Check if the content is not empty
            if k in ["Experience", "Projects"]:
                # For structured sections, ensure individual entries have content
                filtered_entries = [entry for entry in v if entry.get("title") or entry.get("bullets")]
                if filtered_entries:
                    cleaned_sections[k] = filtered_entries
            else:
                # For string-list sections, ensure content is not empty
                filtered_content = [item for item in v if item.strip()]
                if filtered_content:
                    cleaned_sections[k] = filtered_content
    
    # If "Header/Summary" is present but empty after cleaning, remove it
    if "Header/Summary" in cleaned_sections and not cleaned_sections["Header/Summary"]:
        del cleaned_sections["Header/Summary"]

    return cleaned_sections


def batch_refine_resume_gemini(sections, keywords, jd_keywords, gemini_api_key, language="english", position_title="Desired Position"):
    """Refines resume sections using Gemini API with improved prompts."""
    # `sections` should now contain structured data for Experience/Projects from parse_resume_sections
    resume_json = json.dumps(sections, indent=2)
    model = genai.GenerativeModel("gemini-1.5-flash")

    lang_instruction = "Respond in French" if language == "french" else "Respond in English"

    prompt = f"""You are an expert resume writer. Refine the following resume sections to be ATS-friendly and tailored to the job description.

CRITICAL REQUIREMENTS:
1. {lang_instruction} throughout the entire response
2. Preserve ALL original factual information (dates, companies, roles, project names, education details). DO NOT OMIT OR EMPTY EXISTING SECTIONS.
3. Bold relevant keywords using <b></b> HTML tags: {', '.join(keywords + jd_keywords)}
4. Maximum 2 pages total length - be extremely concise
5. For Experience/Projects: Maximum 3-4 bullet points per entry, starting with action verbs. Ensure each entry has detailed accomplishments.
6. Add a compelling Professional Summary (3-4 sentences) highlighting top skills for '{position_title}'
7. If Experience or Projects sections are genuinely empty in the *input*, create realistic entries based on skills/education with "(inferred)" note. If they are *not* empty, preserve and refine their content.

FORMATTING RULES:
- No bullet prefixes (•, -, *, bullet) - start directly with content
- Bold project names: <b>Project Name</b>
- Quantify achievements wherever possible
- Use action verbs (Led, Developed, Implemented, etc.)

SECTION HANDLING:
- Skills: Focus on job-relevant skills only, bold keywords
- Experience: Provide as an array of JSON objects, each with 'title' (e.g., 'Company – Role | Dates') and 'bullets' (array of strings).
- Projects: Provide as an array of JSON objects, each with 'title' (e.g., '<b>Project Name</b>') and 'bullets' (array of strings).
- Education: Institution, Degree, Year, relevant details (e.g., Institution – Degree | Year)

Return ONLY a valid JSON object where section names are keys. For "Experience" and "Projects" sections, the value should be an array of objects, each with "title" (e.g., "Company – Role | Dates" or "Project Name") and "bullets" (array of strings for bullet points). For other sections, the value should be an array of strings.

Example for Experience entry:
{{
  "Experience": [
    {{
      "title": "Infosys – Senior System Engineer | Jan 2024 - Present",
      "bullets": [
        "Designed robust forms using Spring Boot and Spring Data JPA, managing patient, clinic, and clinician records.",
        "Developed real-time monitoring and alerting using WebSockets and Spring Security, reducing response time by 40%.",
        "Led microservices architecture migration, improving system scalability by 30%."
      ]
    }}
  ]
}}

Example for Project entry:
{{
  "Projects": [
    {{
      "title": "<b>Journal App</b>",
      "bullets": [
        "Developed a secure journaling backend using Spring Boot and JWT-based authentication; integrated external APIs and deployed on Heroku.",
        "Leveraged Kafka and Redis for message brokering and caching; built efficient CRUD endpoints with robust error handling."
      ]
    }}
  ]
}}

Resume Sections: {resume_json}"""

    try:
        response = model.generate_content(
            prompt,
            generation_config=genai.GenerationConfig(
                temperature=0.7,
                max_output_tokens=4000
            ),
            safety_settings={
                'HARM_CATEGORY_HARASSMENT': 'BLOCK_NONE',
                'HARM_CATEGORY_HATE_SPEECH': 'BLOCK_NONE',
                'HARM_CATEGORY_SEXUALLY_EXPLICIT': 'BLOCK_NONE',
                'HARM_CATEGORY_DANGEROUS_CONTENT': 'BLOCK_NONE',
            }
        )
        json_match = re.search(r'\{[\s\S]*\}', response.text)
        if json_match:
            refined_data = json.loads(json_match.group(0))
            # The post-processing for Experience and Projects is now primarily handled by parse_resume_sections.
            # We keep a minimal check here for robustness if Gemini deviates.
            for section_key in ['Experience', 'Projects']:
                if section_key in refined_data and isinstance(refined_data[section_key], list):
                    # Ensure each item in the list is a dict with 'title' and 'bullets'
                    processed_items = []
                    for item in refined_data[section_key]:
                        if isinstance(item, dict) and "title" in item and "bullets" in item and isinstance(item["bullets"], list):
                            processed_items.append(item)
                        elif isinstance(item, str): # Fallback if Gemini sends it as flat strings
                            lines = [line.strip() for line in item.split('\n') if line.strip()]
                            if lines:
                                title = lines[0]
                                bullets = lines[1:]
                                processed_items.append({"title": title, "bullets": bullets})
                            else: # If string is empty or just whitespace
                                processed_items.append({"title": "", "bullets": []})
                        # Handle other unexpected dict formats by attempting conversion
                        elif isinstance(item, dict):
                             title_candidate = item.get("title") or next(iter(item.keys()), "")
                             bullets_candidate = [str(v) for v in item.values()] if not item.get("bullets") else item["bullets"]
                             if isinstance(bullets_candidate, str): bullets_candidate = [bullets_candidate] # Ensure bullets is a list
                             processed_items.append({"title": title_candidate, "bullets": bullets_candidate})

                    refined_data[section_key] = processed_items
            return refined_data
        return json.loads(response.text)
    except (json.JSONDecodeError, api_exceptions.GoogleAPIError, Exception) as e:
        st.error(f"Error refining resume: {e}")
        return None

def create_modern_resume_pdf(sections, filename, personal_info, photo_path=None):
    """Creates a modern, ATS-friendly resume PDF."""

    def header_footer(canvas, doc, name, contact_info, photo_path):
        canvas.saveState()
        canvas.setFillColor(darkblue)
        canvas.rect(0, letter[1] - 1.2 * inch, letter[0], 1.2 * inch, fill=1)
        canvas.setFillColor('white')
        canvas.setFont("Helvetica-Bold", 22)
        canvas.drawString(0.5 * inch, letter[1] - 0.6 * inch, name.upper())

        canvas.setFont("Helvetica", 10)
        y_pos = letter[1] - 0.9 * inch
        for info in contact_info:
            if info:
                canvas.drawString(0.5 * inch, y_pos, info)
                y_pos -= 0.15 * inch

        if photo_path and os.path.exists(photo_path):
            try:
                canvas.drawImage(photo_path, letter[0] - 1.3 * inch, letter[1] - 1.1 * inch,
                                 width=0.8 * inch, height=0.8 * inch, mask='auto')
            except Exception as e:
                print(f"Photo error: {e}")
        canvas.restoreState()

    doc = SimpleDocTemplate(filename, pagesize=letter,
                            topMargin=1.3 * inch, bottomMargin=0.5 * inch,
                            leftMargin=0.5 * inch, rightMargin=0.5 * inch)

    def _page_template_wrapper(canvas, doc):
        name = personal_info.get('name', 'Your Name')
        contact_info = [personal_info.get('phone', ''), personal_info.get('email', ''),
                        personal_info.get('linkedin', ''), personal_info.get('github', '')]
        contact_info = [info for info in contact_info if info]
        header_footer(canvas, doc, name, contact_info, photo_path)

    doc.onFirstPage = _page_template_wrapper
    doc.onLaterPages = _page_template_wrapper

    styles = getSampleStyleSheet()

    if 'SectionHeader' not in styles:
        styles.add(ParagraphStyle(name='SectionHeader', fontSize=12, fontName='Helvetica-Bold',
                                  spaceAfter=6, textColor=darkblue))

    body_text_style = styles['BodyText']
    body_text_style.fontSize = 10
    body_text_style.fontName = 'Helvetica'
    body_text_style.spaceAfter = 4
    body_text_style.alignment = TA_JUSTIFY

    if 'BulletText' not in styles:
        styles.add(ParagraphStyle(name='BulletText', fontSize=10, fontName='Helvetica',
                                  spaceAfter=3, leftIndent=0.25 * inch, firstLineIndent=-0.25 * inch))

    story = []

    section_order = [
        "Professional Summary", "Summary", "Objective",
        "Technical Skills", "Skills", "Core Competencies",
        "Experience", "Work Experience", "Professional Experience",
        "Projects", "Education", "Certifications", "Awards"
    ]

    processed = set()

    for section_name in section_order:
        matching_key = None
        for key in sections.keys():
            if key.lower() == section_name.lower() or \
               re.search(r'\b' + re.escape(section_name.lower()) + r'\b', key.lower()):
                matching_key = key
                break

        if matching_key and matching_key not in processed and sections[matching_key]:
            story.append(Paragraph(f"<b>{matching_key.upper()}</b>", styles['SectionHeader']))
            story.append(Spacer(1, 0.1 * inch))

            for item in sections[matching_key]:
                if isinstance(item, dict):  # e.g., Experience/Projects with title and bullets
                    title = item.get("title", "")
                    bullets = item.get("bullets", [])
                    if title:
                        story.append(Paragraph(f"<b>{title}</b>", styles['BodyText']))
                    for bullet in bullets:
                        bullet = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', bullet)
                        bullet = re.sub(r'^(?:bullet\s*|•\s*|[-*]\s*)', '', bullet, flags=re.IGNORECASE).strip()
                        if bullet: # Only add if bullet content exists after stripping
                            story.append(Paragraph(f"• {bullet}", styles['BulletText']))
                elif isinstance(item, str) and item.strip():
                    item = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', item)
                    item = re.sub(r'^(?:bullet\s*|•\s*|[-*]\s*)', '', item, flags=re.IGNORECASE).strip()
                    story.append(Paragraph(f"• {item}", styles['BulletText']))

            story.append(Spacer(1, 0.15 * inch))
            processed.add(matching_key)

    for key, content in sections.items():
        if key not in processed and content:
            story.append(Paragraph(f"<b>{key.upper()}</b>", styles['SectionHeader']))
            story.append(Spacer(1, 0.1 * inch))

            for line in content:
                # This block handles sections that are simple lists of strings (e.g., Skills, Certifications)
                if isinstance(line, str) and line.strip():
                    line_html = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', line)
                    line_html = re.sub(r'^(?:bullet\s*|•\s*|[-*]\s*)', '', line_html, flags=re.IGNORECASE).strip()
                    story.append(Paragraph(f"• {line_html}", styles['BulletText']))
                # If a section like Experience or Projects somehow falls here as a dict list (shouldn't if handled above)
                elif isinstance(line, dict):
                    title = line.get("title", "")
                    bullets = line.get("bullets", [])
                    if title:
                        story.append(Paragraph(f"<b>{title}</b>", styles['BodyText']))
                    for bullet in bullets:
                        bullet = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', bullet)
                        bullet = re.sub(r'^(?:bullet\s*|•\s*|[-*]\s*)', '', bullet, flags=re.IGNORECASE).strip()
                        if bullet:
                            story.append(Paragraph(f"• {bullet}", styles['BulletText']))


            story.append(Spacer(1, 0.15 * inch))

    try:
        doc.build(story)
    except Exception as e:
        st.error(f"Error creating resume PDF: {e}")

def generate_cover_letter_content(google_gemini_api_key, job_description, resume_sections,
                                personal_info, company_name, recruiter_name, position_title, language="english"):
    """Generates cover letter content using Gemini API."""
    model = genai.GenerativeModel("gemini-1.5-flash")

    def format_resume_for_prompt(sections: dict) -> str:
            output = ""
            for section, entries in sections.items():
                output += f"\n### {section}\n"
                for item in entries:
                    if isinstance(item, dict):  # For structured sections
                        title = item.get("title", "")
                        bullets = item.get("bullets", [])
                        if title:
                            output += f"{title}\n"
                        if isinstance(bullets, list):
                            for bullet in bullets:
                                output += f"- {bullet}\n"
                    elif isinstance(item, str):
                        output += f"- {item}\n"
            return output

    resume_summary = format_resume_for_prompt(resume_sections)
    lang_instruction = "Respond in French" if language == "french" else "Respond in English"

    prompt = f"""You are an expert cover letter writer. Create a professional, compelling cover letter in {language}.

REQUIREMENTS:
1. {lang_instruction} throughout
2. Address to "{recruiter_name}" (use "Madame, Monsieur" if French and generic, "Dear Sir or Madam" if English)
3. Position: {position_title} at {company_name}
4. Bold relevant keywords using <b></b> tags
5. Include quantifiable achievements
6. Professional yet enthusiastic tone
7. Maximum 400 words for main content

STRUCTURE (return as JSON):
{{
  "opening": "Opening paragraph expressing interest and company knowledge",
  "body_paragraphs": [
    "Paragraph 1: Relevant experience and skills match",
    "Paragraph 2: Specific achievements and value proposition"
  ],
  "achievements": [
    "Achievement 1 with metrics",
    "Achievement 2 with metrics"
  ],
  "closing": "Professional closing expressing interview interest"
}}

JOB DESCRIPTION: {job_description[:1000]}
RESUME: {resume_summary[:1500]}
PERSONAL INFO: {json.dumps(personal_info)}"""

    try:
        response = model.generate_content(
            prompt,
            generation_config=genai.GenerationConfig(
                temperature=0.7,
                max_output_tokens=1000
            ),
            safety_settings={
                'HARM_CATEGORY_HARASSMENT': 'BLOCK_NONE',
                'HARM_CATEGORY_HATE_SPEECH': 'BLOCK_NONE',
                'HARM_CATEGORY_SEXUALLY_EXPLICIT': 'BLOCK_NONE',
                'HARM_CATEGORY_DANGEROUS_CONTENT': 'BLOCK_NONE',
            }
        )
        json_match = re.search(r'\{[\s\S]*\}', response.text)
        if json_match:
            return json.loads(json_match.group(0))
        return json.loads(response.text)
    except (json.JSONDecodeError, api_exceptions.GoogleAPIError, Exception) as e:
        st.error(f"Error generating cover letter: {e}")
        return None

def create_cover_letter_pdf(filename, personal_info, company_info, position_title,
                          cover_letter_content, language="english"):
    """Creates a professional cover letter PDF."""
    doc = SimpleDocTemplate(filename, pagesize=letter, topMargin=0.75*inch,
                          bottomMargin=0.75*inch, leftMargin=0.75*inch, rightMargin=0.75*inch)

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Address', fontSize=10, fontName='Helvetica', alignment=TA_RIGHT))
    styles.add(ParagraphStyle(name='Date', fontSize=11, fontName='Helvetica', spaceAfter=12))
    styles.add(ParagraphStyle(name='Recipient', fontSize=11, fontName='Helvetica', spaceAfter=12))
    styles.add(ParagraphStyle(name='Subject', fontSize=12, fontName='Helvetica-Bold',
                             spaceAfter=12, alignment=TA_CENTER))
    styles.add(ParagraphStyle(name='Body', fontSize=11, fontName='Helvetica',
                             spaceAfter=12, alignment=TA_JUSTIFY))

    bullet_style = styles['Bullet']
    bullet_style.fontSize = 11
    bullet_style.fontName = 'Helvetica'
    bullet_style.spaceAfter = 6
    bullet_style.leftIndent = 0.25 * inch
    bullet_style.firstLineIndent = -0.25 * inch


    story = []

    sender_lines = [
        personal_info.get('name', ''),
        personal_info.get('address', ''),
        personal_info.get('city', ''),
        f"{personal_info.get('phone', '')} | {personal_info.get('email', '')}"
    ]
    sender_text = "<br/>".join([line for line in sender_lines if line])
    story.append(Paragraph(sender_text, styles['Address']))
    story.append(Spacer(1, 0.3*inch))

    story.append(Paragraph(company_info['date'], styles['Date']))

    recipient_text = f"<b>{company_info['recruiter']}</b><br/>{company_info['company']}<br/>{company_info['company_city']}"
    story.append(Paragraph(recipient_text, styles['Recipient']))

    if language == "french":
        subject = f"<b>Objet : Candidature pour le poste de {position_title}</b>"
        salutation = "Madame, Monsieur," if company_info['recruiter'].lower() in ['hiring manager', 'recruiter'] else f"Cher/Chère {company_info['recruiter']},"
    else:
        subject = f"<b>Subject: Application for the Position of {position_title}</b>"
        salutation = "Dear Sir or Madam," if company_info['recruiter'].lower() in ['hiring manager', 'recruiter'] else f"Dear {company_info['recruiter']},"

    story.append(Paragraph(subject, styles['Subject']))
    story.append(Paragraph(salutation, styles['Body']))

    if cover_letter_content:
        story.append(Paragraph(cover_letter_content.get('opening', ''), styles['Body']))

        for para in cover_letter_content.get('body_paragraphs', []):
            story.append(Paragraph(para, styles['Body']))

        if cover_letter_content.get('achievements'):
            achievements_header = "Mes principales réalisations :" if language == "french" else "Key Achievements:"
            story.append(Paragraph(f"<b>{achievements_header}</b>", styles['Body']))
            for achievement in cover_letter_content['achievements']:
                story.append(Paragraph(f"• {achievement}", styles['Bullet']))

        story.append(Paragraph(cover_letter_content.get('closing', ''), styles['Body']))

    closing_phrase = "Cordialement," if language == "french" else "Sincerely,"
    story.append(Paragraph(closing_phrase, styles['Body']))
    story.append(Spacer(1, 0.2*inch))
    story.append(Paragraph(f"<b>{personal_info.get('name', '')}</b>", styles['Body']))

    enclosure_text = "Pièce jointe : Dossier de candidature" if language == "french" else "Enclosure: Application file"
    story.append(Spacer(1, 0.1*inch))
    story.append(Paragraph(enclosure_text, styles['Body']))

    try:
        doc.build(story)
    except Exception as e:
        st.error(f"Error creating cover letter PDF: {e}")

# --- Streamlit Application ---

st.set_page_config(page_title="AI Resume & Cover Letter Generator", layout="wide")

st.title("🚀 AI Resume & Cover Letter Generator")
st.markdown("Upload your resume, paste a job description, and generate tailored applications in English or French!")

# Sidebar
st.sidebar.header("Configuration")
openai_api_key_placeholder = st.sidebar.text_input("OpenAI API Key (Not used for CL anymore, for dev reference)", type="password", key="openai_api_key_placeholder")
google_gemini_api_key = st.sidebar.text_input("Enter your Google Gemini API Key (for both Resume & Cover Letter)", type="password", key="gemini_api_key_input")

if not google_gemini_api_key:
    st.sidebar.warning("Please enter your Google Gemini API Key to use the application.")
else:
    genai.configure(api_key=google_gemini_api_key)

tab1, tab2 = st.tabs(["📄 Resume Refinement", "📝 Cover Letter Generation"])

with tab1:
    st.header("Resume Refinement")

    col1, col2 = st.columns(2)
    with col1:
        resume_file = st.file_uploader("Upload Resume (PDF, DOCX)", type=["pdf", "docx"], key="resume_uploader")
    with col2:
        photo_file = st.file_uploader("Upload Photo (Optional)", type=["png", "jpg", "jpeg"], key="photo_uploader")

    resume_content = ""
    if resume_file:
        with tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(resume_file.name)[1]) as tmp_file:
            tmp_file.write(resume_file.read())
            temp_file_path = tmp_file.name

        resume_content = extract_text(temp_file_path)
        os.remove(temp_file_path)

        if resume_content:
            st.success("Resume uploaded and text extracted successfully!")

            st.subheader("Personal Information (Edit as needed)")
            extracted_info = extract_personal_details(resume_content)

            col1, col2 = st.columns(2)
            with col1:
                name = st.text_input("Full Name", value=extracted_info.get('name', ''), key="name_input")
                email = st.text_input("Email", value=extracted_info.get('email', ''), key="email_input")
                phone = st.text_input("Phone", value=extracted_info.get('phone', ''), key="phone_input")
            with col2:
                linkedin = st.text_input("LinkedIn URL", value=extracted_info.get('linkedin', ''), key="linkedin_input")
                github = st.text_input("GitHub URL", value=extracted_info.get('github', ''), key="github_input")
                address = st.text_input("Address", value=extracted_info.get('address', ''), key="address_input")
                city = st.text_input("City", value=extracted_info.get('city', ''), key="city_input")

            personal_info = {
                'name': name, 'email': email, 'phone': phone,
                'linkedin': linkedin, 'github': github, 'address': address, 'city': city
            }

            st.subheader("Job Description")
            jd_method = st.radio("Input method:", ("Paste Text", "Enter URL"), key="jd_method_radio")

            if jd_method == "Paste Text":
                job_description = st.text_area("Paste job description here:", height=200, key="jd_text_area")
            else:
                jd_url = st.text_input("Job description URL:", key="jd_url_input")
                job_description = fetch_job_description(jd_url) if jd_url else ""

            position_title = st.text_input("Position Title:", "Software Engineer", key="position_title_tab1")

            if st.button("🚀 Refine Resume", type="primary", key="refine_resume_button"):
                if not google_gemini_api_key:
                    st.error("Please enter your Google Gemini API Key in the sidebar to refine the resume.")
                elif job_description:
                    with st.spinner("Refining your resume..."):
                        resume_lang = detect_language(resume_content)
                        jd_lang = detect_language(job_description)
                        final_lang = resume_lang if resume_lang == jd_lang else 'english'

                        resume_keywords = extract_keywords(resume_content, language=final_lang)
                        jd_keywords = extract_keywords(job_description, language=final_lang)

                        sections = parse_resume_sections(resume_content)

                        st.subheader("DEBUG: Parsed Sections from your Resume")
                        st.json(sections)

                        refined_sections = batch_refine_resume_gemini(
                            sections, resume_keywords, jd_keywords,
                            google_gemini_api_key, final_lang, position_title
                        )

                        if refined_sections:
                            st.session_state.refined_sections = refined_sections
                            st.session_state.personal_info = personal_info
                            st.session_state.job_description = job_description
                            st.session_state.language = final_lang
                            st.session_state.position_title = position_title

                            st.subheader("DEBUG: Refined Sections from Gemini")
                            st.json(refined_sections)

                            temp_pdf = os.path.join(tempfile.gettempdir(), "refined_resume.pdf")

                            photo_path = None
                            if photo_file:
                                with tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(photo_file.name)[1]) as tmp_photo:
                                    tmp_photo.write(photo_file.read())
                                    photo_path = tmp_photo.name

                            create_modern_resume_pdf(refined_sections, temp_pdf, personal_info, photo_path)

                            if os.path.exists(temp_pdf):
                                with open(temp_pdf, "rb") as pdf_file:
                                    st.download_button(
                                        "📥 Download Refined Resume PDF",
                                        pdf_file.read(),
                                        "refined_resume.pdf",
                                        "application/pdf",
                                        key="resume_download"
                                    )
                                st.success("✅ Resume refined and PDF generated successfully!")
                                os.remove(temp_pdf)
                                if photo_path and os.path.exists(photo_path):
                                    os.remove(photo_path)

with tab2:
    st.header("Cover Letter Generation")

    if 'refined_sections' not in st.session_state:
        st.info("Please refine your resume in the 'Resume Refinement' tab first.")
    else:
        st.subheader("Company and Recruiter Details")
        company_name = st.text_input("Company Name:", key="company_name_cl")
        recruiter_name = st.text_input("Recruiter Name (e.g., 'Hiring Manager' or specific name):", key="recruiter_name_cl")
        company_city = st.text_input("Company City:", key="company_city_cl")
        cl_position_title = st.text_input("Position Title (for Cover Letter):", value=st.session_state.get('position_title', 'Software Engineer'), key="position_title_cl")
        cl_date = datetime.date.today().strftime("%B %d, %Y")

        company_info = {
            'company': company_name,
            'recruiter': recruiter_name,
            'company_city': company_city,
            'date': cl_date
        }

        if st.button("✍️ Generate Cover Letter", type="primary", key="generate_cl_button"):
            if not google_gemini_api_key:
                st.error("Please enter your Google Gemini API Key in the sidebar to generate the cover letter.")
            elif not company_name or not recruiter_name or not cl_position_title:
                st.error("Please fill in all company and recruiter details.")
            else:
                with st.spinner("Generating cover letter..."):
                    cover_letter_content = generate_cover_letter_content(
                        google_gemini_api_key,
                        st.session_state.job_description,
                        st.session_state.refined_sections,
                        st.session_state.personal_info,
                        company_info['company'],
                        company_info['recruiter'],
                        cl_position_title,
                        st.session_state.language
                    )

                    if cover_letter_content:
                        st.session_state.cover_letter_content = cover_letter_content
                        st.subheader("DEBUG: Generated Cover Letter Content from Gemini")
                        st.json(cover_letter_content)

                        temp_cl_pdf = os.path.join(tempfile.gettempdir(), "cover_letter.pdf")
                        create_cover_letter_pdf(
                            temp_cl_pdf,
                            st.session_state.personal_info,
                            company_info,
                            cl_position_title,
                            cover_letter_content,
                            st.session_state.language
                        )

                        if os.path.exists(temp_cl_pdf):
                            with open(temp_cl_pdf, "rb") as pdf_file:
                                st.download_button(
                                    "📥 Download Cover Letter PDF",
                                    pdf_file.read(),
                                    "cover_letter.pdf",
                                    "application/pdf",
                                    key="cl_download"
                                )
                            st.success("✅ Cover letter generated and PDF created successfully!")
                            os.remove(temp_cl_pdf)